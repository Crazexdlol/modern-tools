"""One-off: convert the admin claim-workflow guide (Markdown) to a nicely
formatted .docx. Lightweight custom parser — handles headings (# ## ###),
pipe-delimited tables, bulleted + numbered lists, bold **x**, italic *x*,
inline `code`, and horizontal rules (---). Not a general markdown engine;
written specifically for the structure of this one document.
"""

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm, Inches


SRC = Path(r"C:\XSFA\Modern Tools Official\docs\claim-workflow-admin-guide.md")
OUT = Path(r"C:\Users\Owner\Downloads\Claim Workflow — Admin Onboarding Guide.docx")


# ---- Styling helpers --------------------------------------------------------

def set_cell_shading(cell, hex_fill: str) -> None:
    """Set a table cell's background color (hex without the leading #)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_borders(cell, hex_color: str = "DDDDDD") -> None:
    """Thin uniform borders on all four sides of a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), hex_color)
        borders.append(b)
    tc_pr.append(borders)


def add_horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


# ---- Inline-markup renderer (bold / italic / inline code) ------------------

TOKEN_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)"
)

def add_inline_runs(paragraph, text: str, base_font: str = "Calibri", base_size: int = 11, base_color=None) -> None:
    """Split a text run into bold / italic / code segments and attach to the paragraph."""
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
        else:
            run = paragraph.add_run(part)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            if base_color is not None:
                run.font.color.rgb = base_color


# ---- Document setup --------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base body style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    # H1 (page title)
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(22)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(10)

    # H2 (section)
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(16)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)

    # H3 (subsection)
    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x37, 0x47, 0x6B)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)

    return doc


# ---- Markdown parser (purpose-built for this file) -------------------------

TABLE_SEP_RE = re.compile(r"^\s*\|?\s*(:?-+:?\s*\|)+\s*:?-+:?\s*\|?\s*$")
TABLE_ROW_RE = re.compile(r"^\s*\|.+\|\s*$")
ORDERED_RE = re.compile(r"^\s*(\d+)\.\s+(.+)$")
UNORDERED_RE = re.compile(r"^\s*[-*]\s+(.+)$")


def split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [c.strip() for c in stripped.split("|")]


def render_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = True

    # Header row
    for i, cell_text in enumerate(header):
        cell = tbl.rows[0].cells[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = ""  # clear default paragraph
        p = cell.paragraphs[0]
        add_inline_runs(p, cell_text)
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "1F2A44")  # dark navy header
        set_cell_borders(cell, "1F2A44")

    # Body rows
    for r, row in enumerate(rows, start=1):
        for i, cell_text in enumerate(row):
            cell = tbl.rows[r].cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_runs(p, cell_text)
            set_cell_borders(cell, "DDDDDD")
            # Light alternating row shading for readability
            if r % 2 == 0:
                set_cell_shading(cell, "F6F7FB")

    # Little spacing after the table
    doc.add_paragraph()


def render_list(doc: Document, items: list[tuple[str, str]]) -> None:
    """items: list of (style_name, text) — style_name = 'List Number' or 'List Bullet'"""
    for style, text in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        add_inline_runs(p, text)


def convert(md: str, doc: Document) -> None:
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            add_horizontal_rule(doc)
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
            i += 1
            continue

        # Table — look ahead for header + separator
        if TABLE_ROW_RE.match(line) and i + 1 < len(lines) and TABLE_SEP_RE.match(lines[i + 1]):
            header = split_table_row(line)
            rows = []
            j = i + 2
            while j < len(lines) and TABLE_ROW_RE.match(lines[j]):
                rows.append(split_table_row(lines[j]))
                j += 1
            render_table(doc, header, rows)
            i = j
            continue

        # List block (collect consecutive list lines)
        if ORDERED_RE.match(line) or UNORDERED_RE.match(line):
            items = []
            while i < len(lines):
                cur = lines[i]
                m_ord = ORDERED_RE.match(cur)
                m_un = UNORDERED_RE.match(cur)
                if m_ord:
                    items.append(("List Number", m_ord.group(2)))
                elif m_un:
                    items.append(("List Bullet", m_un.group(1)))
                elif cur.startswith(" " * 3) and items:
                    # Continuation of the previous item (e.g. a sub-line)
                    items[-1] = (items[-1][0], items[-1][1] + " " + cur.strip())
                else:
                    break
                i += 1
            render_list(doc, items)
            continue

        # Normal paragraph (join consecutive non-blank non-structural lines)
        buf = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            if (
                not nxt.strip()
                or nxt.strip().startswith("#")
                or nxt.strip() == "---"
                or TABLE_ROW_RE.match(nxt)
                or ORDERED_RE.match(nxt)
                or UNORDERED_RE.match(nxt)
            ):
                break
            buf.append(nxt)
            j += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(line.strip() for line in buf))
        i = j


def main() -> int:
    if not SRC.exists():
        print(f"Source not found: {SRC}", file=sys.stderr)
        return 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    md = SRC.read_text(encoding="utf-8")
    doc = build_document()
    convert(md, doc)
    doc.save(OUT)
    print(f"Wrote: {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
